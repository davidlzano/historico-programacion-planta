"""
Genera los respaldos diarios sobre los que trabaja el consolidador.

Reproduce a propósito el problema que hace difícil este trabajo: **el layout
cambia entre archivos**. De un día a otro las columnas se mueven, cambian de
nombre, aparecen nuevas y a veces una hoja queda ilegible.

Eso no es un defecto del archivo origen: es un libro de Excel que edita a mano
el área de producción todos los días. Esperar que la estructura se mantenga
estable es la suposición que rompe cualquier carga que lea por posición.

Uso:
    python generar_snapshots.py
"""

import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl

DIR_SNAPSHOTS = Path("snapshots")
DIAS = 20
SEMILLA = 11

CENTROS = ["IMPRESION 1", "IMPRESION 2", "TROQUELADO", "CONV.", "EMP MANUAL"]

CLIENTES = ["ALFA", "BETA", "GAMMA", "DELTA", "OMEGA", "SIGMA"]
PRODUCTOS = ["CAJA PLEGADIZA", "ESTUCHE", "MICROCORRUGADO", "DISPLAY"]

# El mismo campo aparece con nombres distintos según quién editó la hoja.
ALIAS = {
    "op": ["# OP", "OP", "N OP"],
    "cliente": ["CLIENTE"],
    "referencia": ["REFERENCIA", "REFER", "CODIGO"],
    "descripcion": ["DESCRIPCION", "DESCRIPCIÓN"],
    "cantidad": ["TIROS", "CANTIDAD", "CANTIDAD PEDIDA"],
    "inicio": ["DIA Y HORA INICIO"],
    "fin": ["DIA Y HORA FIN"],
    "alistamiento": ["TIEMPO ALISTAMIENTO", "TIEMPO ALISTO"],
    "observaciones": ["OBSERVACIONES", "OBS"],
}

# Filas que no son producción sino eventos de calendario o de máquina.
EVENTOS = [
    "MANTENIMIENTO PREVENTIVO",
    "LAVADO DE MAQUINA",
    "INICIO DE TURNO",
    "DOMINGO",
    "FESTIVO",
    "PARO POR FALTA DE MATERIAL",
]


def _programa_del_dia(rnd: random.Random, centro: str, dia: date) -> list[dict]:
    """Programación de un centro de trabajo para un día."""
    filas = []
    hora = 6
    for _ in range(rnd.randint(3, 7)):
        if rnd.random() < 0.15:
            filas.append({"op": "", "descripcion": rnd.choice(EVENTOS), "cantidad": ""})
            hora += rnd.randint(1, 3)
            continue

        cliente = rnd.choice(CLIENTES)
        duracion = rnd.randint(2, 6)
        filas.append(
            {
                "op": f"OP-{rnd.randint(2600, 2799)}",
                "cliente": cliente,
                "referencia": f"{cliente[:3]}-{rnd.randint(100, 999)}",
                "descripcion": f"{rnd.choice(PRODUCTOS)} {cliente}",
                "cantidad": rnd.randrange(1000, 40000, 500),
                "inicio": f"{dia.isoformat()} {hora:02d}:00",
                "fin": f"{dia.isoformat()} {min(hora + duracion, 23):02d}:00",
                "alistamiento": round(rnd.uniform(0.5, 2.5), 1),
                "observaciones": rnd.choice(["", "", "", "URGENTE", "REPROCESO"]),
            }
        )
        hora += duracion
        if hora >= 22:
            hora = 6
    return filas


def _escribir_hoja(ws, filas: list[dict], rnd: random.Random, dia_indice: int) -> None:
    """
    Escribe una hoja con layout variable.

    Tres fuentes de variación, todas observadas en el archivo real:
      - el encabezado no siempre arranca en la fila 1
      - el orden de las columnas cambia
      - el mismo campo aparece con nombres distintos
    """
    campos = list(ALIAS.keys())
    rnd.shuffle(campos)

    # A partir de cierto día se agrega una columna nueva al libro.
    if dia_indice >= 12 and "medida" not in campos:
        campos.append("medida")

    fila_encabezado = rnd.choice([1, 2, 3])
    for _ in range(fila_encabezado - 1):
        ws.append([])

    encabezados = [
        rnd.choice(ALIAS[c]) if c in ALIAS else "MEDIDA" for c in campos
    ]
    ws.append(encabezados)

    for fila in filas:
        ws.append([fila.get(c, "") for c in campos])


def main() -> None:
    rnd = random.Random(SEMILLA)
    DIR_SNAPSHOTS.mkdir(exist_ok=True)
    for viejo in DIR_SNAPSHOTS.glob("*.xlsx"):
        viejo.unlink()

    hoy = date.today()
    # Estado que evoluciona día a día: parte de él se mantiene, parte cambia.
    estado: dict[str, list[dict]] = {}

    for i in range(DIAS):
        dia = hoy - timedelta(days=DIAS - 1 - i)
        # Fin de semana sin respaldo, como en la operación real.
        if dia.weekday() >= 5:
            continue

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for centro in CENTROS:
            # El primer día se genera todo; después, la mayoría de la
            # programación se conserva y solo una parte cambia. Eso es lo que
            # el modo de solo-cambios debe aprovechar.
            if centro not in estado or rnd.random() < 0.08:
                estado[centro] = _programa_del_dia(rnd, centro, dia)
            else:
                for fila in estado[centro]:
                    if fila.get("op") and rnd.random() < 0.04:
                        fila["cantidad"] = rnd.randrange(1000, 40000, 500)

            # Una hoja ocasionalmente queda ilegible: contraseña, corrupción,
            # o alguien la dejó en un estado que openpyxl no puede abrir.
            # El consolidador debe registrarlo como aviso, no como datos
            # borrados.
            if rnd.random() < 0.06:
                ws = wb.create_sheet(centro)
                ws.append(["HOJA NO DISPONIBLE"])
                continue

            ws = wb.create_sheet(centro)
            _escribir_hoja(ws, estado[centro], rnd, i)

        nombre = f"PROGRAMACION {dia.strftime('%d%m%Y')}.xlsx"
        wb.save(DIR_SNAPSHOTS / nombre)

    def _fecha(p):
        return (p.stem[-4:], p.stem[-6:-4], p.stem[-8:-6])

    archivos = sorted(DIR_SNAPSHOTS.glob("*.xlsx"), key=_fecha)
    print(f"{len(archivos)} respaldos generados en {DIR_SNAPSHOTS.resolve()}")
    print(f"Del {archivos[0].name} al {archivos[-1].name}")
    print("\nCada archivo tiene layout distinto: columnas movidas, renombradas,")
    print("encabezados en filas diferentes y alguna hoja ilegible.")


if __name__ == "__main__":
    main()
